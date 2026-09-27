from __future__ import annotations

import csv
import json
import re
from pathlib import Path


def _key(name: str) -> str:
    return re.sub(r"[^a-z0-9]", "", name.lower())


def _num(value):
    try:
        if value is None or str(value).strip() == "":
            return None
        return float(value)
    except (TypeError, ValueError):
        return None


def _read_csv(path: str | Path) -> list[dict]:
    with Path(path).open(encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def _read_lineup_xlsx(path: str | Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Install spreadsheet support with: pip install openpyxl") from exc
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(x).strip() if x is not None else "" for x in rows[0]]
    result = []
    for values in rows[1:]:
        row = dict(zip(headers, values))
        label = str(values[0] or "").strip()
        m = re.match(r"^(.*?) \((.*?) - (.*?)\)$", label)
        if not m:
            continue
        name, team, positions = m.groups()
        row.update({"Player": name, "Team": team, "Positions": positions})
        result.append(row)
    return result


def build_consensus_snapshot(
    fantasypros: str | Path,
    lineup_experts: str | Path,
    output: str | Path,
    *,
    adp: str | Path | None = None,
    ecr: str | Path | None = None,
) -> Path:
    fp = [r for r in _read_csv(fantasypros) if r.get("Player")]
    le = _read_lineup_xlsx(lineup_experts)
    le_by = {_key(r["Player"]): r for r in le}
    adp_by = {_key(r["Player"]): r for r in _read_csv(adp) if r.get("Player")} if adp else {}
    ecr_rows = _read_csv(ecr) if ecr else []
    ecr_by = {}
    for r in ecr_rows:
        label = r.get("PLAYER NAME") or ""
        name = re.sub(r"\s*\([^)]*\)\s*$", "", label)
        if name and _key(name) not in ecr_by:
            ecr_by[_key(name)] = r

    mapping = {"PTS": "PTS", "REB": "RB", "AST": "AST", "BLK": "BLK",
               "STL": "STL", "3PM": "FG3", "TO": "TO"}
    players = []
    for r in fp:
        k = _key(r["Player"])
        other = le_by.get(k)
        stats = {}
        diffs = []
        for cat, le_col in mapping.items():
            a = _num(r.get(cat))
            b = _num(other.get(le_col)) if other else None
            if a is None:
                continue
            stats[cat] = round((a + b) / 2, 3) if b is not None else a
            if b is not None:
                diffs.append(abs(a - b) / max(abs(a), abs(b), 0.5))
        stats["FG%"] = _num(r.get("FG%")) or 0.0
        stats["FT%"] = _num(r.get("FT%")) or 0.0

        market = adp_by.get(k, {})
        expert = ecr_by.get(k, {})
        metadata = {
            "_YAHOO_ADP": _num(market.get("Yahoo")),
            "_ESPN_ADP": _num(market.get("ESPN")),
            "_CONSENSUS_ADP": _num(market.get("AVG")),
            "_ECR": _num(expert.get("RK")),
            "_ECR_SD": _num(expert.get("STD.DEV")),
            "_PROJ_DISAGREEMENT": round(sum(diffs) / len(diffs), 4) if diffs else 0.0,
            "_FP_GP": _num(r.get("GP")),
            "_LE_GP": _num(other.get("GP")) if other else None,
        }
        stats.update({a: b for a, b in metadata.items() if b is not None})
        gps = [x for x in (metadata["_FP_GP"], metadata["_LE_GP"]) if x is not None]
        if gps:
            stats["_CONSENSUS_GP"] = round(sum(gps) / len(gps), 1)

        players.append({
            "player_id": k,
            "name": r["Player"],
            "positions": [p.strip() for p in str(r.get("Positions", "")).split(",") if p.strip()],
            "team": r.get("Team", ""),
            "games_remaining": 0,
            "stats": stats,
            "minutes": _num(r.get("MIN")),
        })

    payload = {
        "as_of": "2026-09-27",
        "settings": {
            "league_id": "elms-2026",
            "name": "ELMS 12-Team 9-Cat",
            "categories": ["PTS", "REB", "AST", "3PM", "STL", "BLK", "FG%", "FT%", "TO"],
            "roster_slots": ["PG", "SG", "G", "SF", "PF", "F", "C", "UTIL", "UTIL", "BN", "BN", "BN", "IL", "IL"],
            "scoring_type": "categories",
        },
        "my_team": {"team_id": "me", "name": "My Team", "players": []},
        "opponents": [],
        "free_agents": players,
    }
    target = Path(output)
    target.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    return target
